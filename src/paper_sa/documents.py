import io
import re
from pathlib import Path
from xml.sax.saxutils import escape

from bs4 import BeautifulSoup
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt
from markdown_it import MarkdownIt
from pypdf import PdfReader
from reportlab.lib import colors
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.cidfonts import UnicodeCIDFont
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from .network import ServiceError

MAX_PDF_BYTES = 60 * 1024 * 1024


def read_pdf(path: Path):
    if path.stat().st_size > MAX_PDF_BYTES:
        raise ServiceError("PDF 超过 60 MB 限制")
    with path.open("rb") as handle:
        if handle.read(5) != b"%PDF-":
            raise ServiceError("文件不是有效 PDF，可能是登录页或验证页面")
    try:
        reader = PdfReader(path)
        if reader.is_encrypted and not reader.decrypt(""):
            raise ServiceError("PDF 受密码保护，请先在本地解密")
        if not 1 <= len(reader.pages) <= 500:
            raise ServiceError("仅支持 1–500 页的 PDF")
        return [(page.extract_text() or "")[:30000] for page in reader.pages]
    except ServiceError:
        raise
    except Exception:
        raise ServiceError("PDF 损坏或无法解析，请尝试重新下载") from None


def pdf_name(paper):
    metadata = paper["metadata"]
    author = (metadata.get("authors") or ["Unknown"])[0].split(" ")[0]
    parts = [
        metadata.get("year") or "Undated",
        metadata.get("journal") or "Paper",
        author,
        metadata.get("title") or "Untitled",
    ]
    name = re.sub(r'[<>:"/\\|?*\x00-\x1f]', "-", "_".join(parts))
    name = re.sub(r"\s+", " ", name).strip(" .")[:140]
    return f"{name}_{paper['id'][:8]}.pdf"


def terms(text):
    words = re.findall(r"[a-z0-9]{2,}", text.lower())
    for segment in re.findall(r"[\u4e00-\u9fff]+", text):
        words.extend(segment[index : index + 2] for index in range(max(1, len(segment) - 1)))
    return set(words)


def paper_context(store, paper, question="", limit=32000, label="P1"):
    metadata = paper["metadata"]
    header = f"[{label}] {metadata.get('title', '')}\nPMID: {paper.get('pmid') or '无'}; DOI: {paper.get('doi') or '无'}\n"
    pages = store.pages(paper["id"])
    if paper["text_status"] != "fulltext" or not pages:
        scope = "扫描件未 OCR，仅使用摘要" if paper["text_status"] == "scan" else "仅摘要，未读取全文"
        abstract = metadata.get("abstract") or "没有摘要或可提取文本，不能得出研究结论。"
        return {
            "text": f"{header}证据范围：{scope}\n[{label} · 摘要]\n{abstract[:limit]}",
            "scope": scope,
            "pages": [],
        }
    chunks = []
    question_terms = terms(question)
    for page in pages:
        for offset in range(0, len(page["text"]), 1600):
            content = page["text"][offset : offset + 1800]
            if content.strip():
                chunks.append(
                    {
                        "page": page["number"],
                        "offset": offset,
                        "text": content,
                        "score": len(question_terms & terms(content)),
                    }
                )
    if question_terms:
        ranked = sorted(chunks, key=lambda item: (-item["score"], item["page"], item["offset"]))
    else:
        ranked = sorted(chunks, key=lambda item: (item["offset"], item["page"]))
    selected, length = [], len(header)
    for chunk in ranked:
        if length + len(chunk["text"]) + 30 <= limit:
            selected.append(chunk)
            length += len(chunk["text"]) + 30
    selected.sort(key=lambda item: (item["page"], item["offset"]))
    selected_pages = sorted({chunk["page"] for chunk in selected})
    scope = f"PDF 文本摘录（{len(selected_pages)}/{paper['page_count']} 页；非完整全文）"
    if len(selected) == len(chunks):
        scope = f"PDF 可提取文本（{paper['page_count']} 页，不含图像和未提取的表格）"
    text = (
        header
        + "证据范围："
        + scope
        + "\n\n"
        + "\n\n".join(f"[{label} · p.{chunk['page']}]\n{chunk['text']}" for chunk in selected)
    )
    return {"text": text, "scope": scope, "pages": selected_pages}


def markdown_blocks(content):
    renderer = MarkdownIt("commonmark", {"html": False}).enable("table")
    soup = BeautifulSoup(renderer.render(content), "html.parser")
    return list(soup.children)


def export_docx(title, content):
    document = Document()
    normal = document.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    for style in ("Title", "Heading 1", "Heading 2", "Heading 3"):
        document.styles[style].font.name = "Calibri"
        document.styles[style].element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    document.add_heading(title, 0)
    for block in markdown_blocks(content):
        if not getattr(block, "name", None):
            continue
        if re.fullmatch(r"h[1-6]", block.name):
            document.add_heading(block.get_text(), min(int(block.name[1]), 3))
        elif block.name in {"ul", "ol"}:
            for item in block.find_all("li", recursive=False):
                document.add_paragraph(
                    item.get_text(" ", strip=True),
                    style="List Bullet" if block.name == "ul" else "List Number",
                )
        elif block.name == "table":
            rows = block.find_all("tr")
            width = max((len(row.find_all(["td", "th"])) for row in rows), default=0)
            if not width:
                continue
            table = document.add_table(rows=0, cols=width)
            table.style = "Light Shading Accent 1"
            for row_index, row in enumerate(rows):
                cells = table.add_row().cells
                for cell, source in zip(cells, row.find_all(["td", "th"])):
                    cell.text = source.get_text(" ", strip=True)
                if row_index == 0:
                    table.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
        else:
            document.add_paragraph(block.get_text("\n", strip=True))
    output = io.BytesIO()
    document.save(output)
    return output.getvalue()


def export_pdf(title, content):
    pdfmetrics.registerFont(UnicodeCIDFont("STSong-Light"))
    styles = getSampleStyleSheet()
    body = ParagraphStyle(
        "CJKBody",
        parent=styles["BodyText"],
        fontName="STSong-Light",
        fontSize=10,
        leading=16,
        wordWrap="CJK",
        spaceAfter=7,
    )
    heading = ParagraphStyle("CJKHeading", parent=body, fontSize=15, leading=21, spaceBefore=12, spaceAfter=9)
    small = ParagraphStyle("CJKSmall", parent=body, fontSize=8, leading=12)
    output = io.BytesIO()
    document = SimpleDocTemplate(
        output,
        pagesize=(210 * mm, 297 * mm),
        rightMargin=18 * mm,
        leftMargin=18 * mm,
        topMargin=18 * mm,
        bottomMargin=18 * mm,
        title=title,
    )
    story = [Paragraph(escape(title), heading), Spacer(1, 5 * mm)]
    for block in markdown_blocks(content):
        if not getattr(block, "name", None):
            continue
        if block.name == "table":
            rows = [
                [
                    Paragraph(escape(cell.get_text(" ", strip=True)), small)
                    for cell in row.find_all(["th", "td"])
                ]
                for row in block.find_all("tr")
            ]
            width = max(map(len, rows), default=0)
            if not width:
                continue
            rows = [row + [""] * (width - len(row)) for row in rows]
            table = Table(
                rows, colWidths=[174 * mm / width] * width, repeatRows=1, splitByRow=1, splitInRow=1
            )
            table.setStyle(
                TableStyle(
                    [
                        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#e9f1ed")),
                        ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#d5dfda")),
                        ("VALIGN", (0, 0), (-1, -1), "TOP"),
                        ("LEFTPADDING", (0, 0), (-1, -1), 6),
                        ("RIGHTPADDING", (0, 0), (-1, -1), 6),
                    ]
                )
            )
            story.extend([table, Spacer(1, 4 * mm)])
        elif block.name in {"ul", "ol"}:
            for index, item in enumerate(block.find_all("li", recursive=False), 1):
                prefix = f"{index}. " if block.name == "ol" else "• "
                story.append(Paragraph(escape(prefix + item.get_text(" ", strip=True)), body))
        else:
            text = block.get_text("\n", strip=True)
            style = heading if re.fullmatch(r"h[1-6]", block.name) else body
            for start in range(0, max(len(text), 1), 5000):
                story.append(Paragraph(escape(text[start : start + 5000]).replace("\n", "<br/>"), style))

    def page_number(canvas, doc):
        canvas.setFont("STSong-Light", 9)
        canvas.drawRightString(192 * mm, 10 * mm, str(doc.page))

    document.build(story, onFirstPage=page_number, onLaterPages=page_number)
    return output.getvalue()
